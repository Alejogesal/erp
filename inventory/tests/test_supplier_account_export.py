"""Descarga en Excel de la cuenta corriente con proveedores: el resumen de todos
y el detalle de movimientos de cada uno."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from inventory.models import Purchase, Supplier, SupplierPayment, Warehouse


class SupplierAccountExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="u", password="x")
        self.client.force_login(self.user)
        self.comun = Warehouse.objects.get(type=Warehouse.WarehouseType.COMUN)
        self.supplier = Supplier.objects.create(name="Aris Norma")
        self.other = Supplier.objects.create(name="Proveedor Sin Movimientos")
        self.purchase = Purchase.objects.create(
            supplier=self.supplier, warehouse=self.comun, total=Decimal("1000.00"), user=self.user
        )
        SupplierPayment.objects.create(supplier=self.supplier, amount=Decimal("300.00"))
        SupplierPayment.objects.create(
            supplier=self.supplier, amount=Decimal("50.00"), kind=SupplierPayment.Kind.ADJUSTMENT
        )

    def _rows(self, resp):
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        rows = list(load_workbook(BytesIO(resp.content)).active.iter_rows(values_only=True))
        return rows[0], rows[1:]

    def test_summary_of_all_suppliers(self):
        resp = self.client.get(reverse("inventory_suppliers"), {"export": "xlsx"})
        header, data = self._rows(resp)
        self.assertEqual(header, ("Proveedor", "Teléfono", "Compras", "Pagos", "Ajustes/Devoluciones", "Saldo CC"))
        by_name = {r[0]: r for r in data}
        self.assertEqual(by_name["Aris Norma"][2:], (1000, 300, 50, 750))
        self.assertEqual(by_name["Proveedor Sin Movimientos"][5], 0)

    def test_summary_adds_period_column_when_filtered(self):
        today = timezone.localdate().isoformat()
        resp = self.client.get(
            reverse("inventory_suppliers"), {"export": "xlsx", "start_date": today, "end_date": today}
        )
        header, data = self._rows(resp)
        self.assertEqual(header[-1], "Comprado (período)")
        self.assertEqual({r[0]: r[-1] for r in data}["Aris Norma"], 1000)

    def test_ledger_of_one_supplier(self):
        resp = self.client.get(
            reverse("inventory_supplier_history", args=[self.supplier.id]), {"export": "xlsx"}
        )
        self.assertIn("cuenta_corriente_aris-norma", resp["Content-Disposition"])
        header, data = self._rows(resp)
        self.assertEqual(header, ("Fecha", "Movimiento", "Detalle", "Debe", "Haber", "Saldo"))
        self.assertEqual(len(data), 3)
        self.assertEqual(data[-1][5], 750)  # saldo final = saldo actual
        self.assertIn(self.purchase.invoice_number, [r[1] for r in data])

    def test_ledger_with_start_date_carries_opening_balance(self):
        start = timezone.localdate() + timedelta(days=1)
        SupplierPayment.objects.create(supplier=self.supplier, amount=Decimal("100.00"), paid_at=start)
        resp = self.client.get(
            reverse("inventory_supplier_history", args=[self.supplier.id]),
            {"export": "xlsx", "start_date": start.isoformat()},
        )
        _, data = self._rows(resp)
        self.assertEqual(data[0][1], "Saldo anterior")
        self.assertEqual(data[0][5], 750)
        self.assertEqual(len(data), 2)
        self.assertEqual(data[1][5], 650)
