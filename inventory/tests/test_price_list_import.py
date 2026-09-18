"""Importar la lista de precios de un proveedor (XLSX): cada proveedor se
respeta tal cual, sin fusionar por parecido de nombre con lo que ya cargó
OTRO proveedor."""

from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from openpyxl import Workbook

from inventory.models import Product, Supplier, SupplierProduct


def _xlsx_upload(rows, filename="lista.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(["Grupo", "Nombre", "Precio", "IVA"])
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        filename,
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class PriceListImportNoFuzzyMatchTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="u", password="x")
        self.client.force_login(self.user)
        self.aris = Supplier.objects.create(name="Aris Norma")
        self.glm = Supplier.objects.create(name="GLM Distribuidora")

    def _import(self, supplier, rows, vat="0"):
        return self.client.post(
            reverse("inventory_suppliers"),
            {
                "action": "import_price_list",
                "price_supplier_id": supplier.id,
                "price_file": _xlsx_upload(rows),
                "price_vat": vat,
            },
        )

    def test_similar_but_different_names_across_suppliers_stay_separate(self):
        # Aris ya tiene cargado "Fidelite Coloracion 8 Rubio Oscuro".
        self._import(self.aris, [["FIDELITE", "Fidelite Coloracion 8 Rubio Oscuro", 100]])
        self.assertEqual(Product.objects.count(), 1)

        # GLM importa "Fidelite Coloracion 8 Rubio" (mismo número "8", y sus
        # palabras son subconjunto de las del producto de Aris): antes, el
        # matcheo difuso (subconjunto de palabras con mismos números) los
        # fusionaba en el mismo Product aunque son tonos distintos. Ahora
        # tienen que quedar separados, cada uno solo vinculado a su proveedor.
        self._import(self.glm, [["FIDELITE", "Fidelite Coloracion 8 Rubio", 90]])

        self.assertEqual(Product.objects.count(), 2)
        aris_product = Product.objects.get(name="Fidelite Coloracion 8 Rubio Oscuro")
        glm_product = Product.objects.get(name="Fidelite Coloracion 8 Rubio")
        self.assertTrue(SupplierProduct.objects.filter(supplier=self.aris, product=aris_product).exists())
        self.assertFalse(SupplierProduct.objects.filter(supplier=self.glm, product=aris_product).exists())
        self.assertTrue(SupplierProduct.objects.filter(supplier=self.glm, product=glm_product).exists())
        self.assertFalse(SupplierProduct.objects.filter(supplier=self.aris, product=glm_product).exists())

    def test_exact_same_name_across_suppliers_still_merges_into_one_product(self):
        # El match EXACTO (mismo nombre normalizado) sigue funcionando: es el
        # caso legítimo de "el mismo producto, dos proveedores lo tienen".
        self._import(self.aris, [["FIDELITE", "Fidelite Coloracion 8", 100]])
        self._import(self.glm, [["FIDELITE", "  fidelite   coloracion 8  ", 90]])

        self.assertEqual(Product.objects.count(), 1)
        product = Product.objects.get()
        self.assertTrue(SupplierProduct.objects.filter(supplier=self.aris, product=product).exists())
        self.assertTrue(SupplierProduct.objects.filter(supplier=self.glm, product=product).exists())

    def test_only_products_of_pinned_supplier_reach_the_price_list(self):
        # Escenario completo reportado: Aris es principal de FIDELITE. Antes,
        # el matcheo difuso podía "colar" un producto de GLM (incluso con
        # precio y todo) bajo el mismo Product de Aris. Ahora, si Aris no
        # tiene ese producto exacto, no aparece en la lista propia.
        self._import(self.aris, [["FIDELITE", "Fidelite Coloracion 8 Rubio Oscuro", 100]])
        self._import(self.glm, [["FIDELITE", "Fidelite Coloracion 8 Rubio", 90]])
        from inventory.models import BrandSupplier

        BrandSupplier.objects.create(group="FIDELITE", supplier=self.aris)
        from inventory import services

        for p in Product.objects.all():
            services.sync_principal_to_cheapest(p)

        url = reverse("inventory_product_prices_download", args=["consumer"])
        resp = self.client.get(url)
        from openpyxl import load_workbook

        ws = load_workbook(BytesIO(resp.content)).active
        nombres = [r[1] for r in ws.iter_rows(min_row=2, values_only=True)]
        self.assertIn("Fidelite Coloracion 8 Rubio Oscuro", nombres)
        self.assertNotIn("Fidelite Coloracion 8 Rubio", nombres)
